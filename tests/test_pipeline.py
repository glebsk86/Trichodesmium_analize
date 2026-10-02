"""Known physical/raster dimensions, failure modes and real CLI exports."""
import csv
import json

import cv2
import numpy as np
from PIL import Image
import pytest
from openpyxl import load_workbook

from trichodesmium.calibration import auto_scale, manual_scale, point_scale
from trichodesmium.cli import main
from trichodesmium.geometry import measure
from trichodesmium.imaging import read_image, read_labels


def ruler(spacing=12,angle=0):
    rgb=np.full((480,640,3),245,np.uint8)
    cv2.line(rgb,(60,240),(580,240),(25,25,25),2)
    for x in range(64,580,spacing): cv2.line(rgb,(x,226),(x,254),(25,25,25),2)
    cv2.line(rgb,(320,50),(320,430),(25,25,25),2)
    for y in range(52,430,spacing): cv2.line(rgb,(306,y),(334,y),(25,25,25),2)
    if angle:
        rgb=cv2.warpAffine(rgb,cv2.getRotationMatrix2D((320,240),angle,1),(640,480),borderValue=(245,245,245))
    return rgb


def rod():
    mask=np.zeros((100,240),bool)
    mask[44:56,30:210]=True
    rgb=np.full((100,240,3),245,np.uint8)
    rgb[mask]=(170,150,70)
    for x in range(42,209,12): rgb[44:56,x]=(60,55,35)
    return mask,rgb


@pytest.mark.parametrize("angle",[0,23,65])
def test_auto_scale_rotated_rulers(angle):
    result=auto_scale(ruler(angle=angle))
    assert result.um_per_px is not None,result.to_dict()
    assert result.um_per_px==pytest.approx(.25,rel=.06)
    assert len(result.axes)==2
    assert result.quality!="высокая"


def test_no_ruler_does_not_invent_scale():
    result=auto_scale(np.full((400,500,3),240,np.uint8))
    assert result.um_per_px is None
    assert "no_regular_ruler_ticks" in result.reasons


def test_repeating_cell_septa_do_not_become_a_ruler():
    _,rgb=rod()
    assert auto_scale(rgb).um_per_px is None
    gray=cv2.cvtColor(rgb,cv2.COLOR_RGB2GRAY)
    assert auto_scale(cv2.cvtColor(gray,cv2.COLOR_GRAY2RGB)).um_per_px is None


@pytest.mark.parametrize("factor",[.75,1.5])
def test_scale_recomputed_after_uniform_resize(factor):
    image=cv2.resize(ruler(),None,fx=factor,fy=factor)
    assert auto_scale(image).um_per_px==pytest.approx(.25/factor,rel=.06)


@pytest.mark.parametrize("angle",[27,45,65])
def test_rotated_filament_length(angle):
    mask=np.zeros((300,300),np.uint8)
    direction=np.array([np.cos(np.deg2rad(angle)),np.sin(np.deg2rad(angle))])
    a,b=np.round(np.array([150,150])+np.array([-80,80])[:,None]*direction).astype(int)
    cv2.line(mask,tuple(a),tuple(b),1,12)
    result=measure(mask.astype(bool),np.full((300,300,3),180,np.uint8),np.zeros_like(mask,bool))
    assert result["length_px"]==pytest.approx(np.linalg.norm(b-a)+12,abs=2)


@pytest.mark.parametrize("bad",[0,-1,float("nan"),float("inf")])
def test_invalid_manual_scale(bad):
    with pytest.raises(ValueError): manual_scale(bad)


def test_reference_points():
    assert point_scale([10,10,70,90],30).um_per_px==pytest.approx(.3)
    with pytest.raises(ValueError): point_scale([1,1,1,1],30)


def test_known_rod_dimensions_and_cell_estimate():
    mask,rgb=rod()
    result=measure(mask,rgb,np.zeros_like(mask))
    assert result["length_px"]==pytest.approx(180,abs=2)
    assert result["width_px"]==pytest.approx(12,abs=.5)
    assert result["cell_length_px"]==pytest.approx(12,abs=.5)
    assert result["cell_count"]==15
    assert result["cell_count_method"]=="estimated_from_length_unverified_septa"


def test_branched_mask_does_not_report_one_chain_length():
    mask=np.zeros((200,200),np.uint8)
    cv2.line(mask,(20,100),(180,100),1,10)
    cv2.line(mask,(100,20),(100,180),1,10)
    result=measure(mask.astype(bool),np.full((200,200,3),180,np.uint8),np.zeros_like(mask,bool))
    assert result["length_px"] is None
    assert "branched_or_closed_mask" in result["flags"]


def test_disconnected_instances_are_rejected():
    mask,rgb=rod()
    mask[10:12,10:12]=True
    result=measure(mask,rgb,np.zeros_like(mask))
    assert result["length_px"] is None


def test_ruler_overlap_excludes_width_samples():
    mask,rgb=rod()
    obstruction=np.zeros_like(mask)
    obstruction[:,110:130]=True
    measured=measure(mask,rgb,obstruction)
    plain=measure(mask,rgb,np.zeros_like(mask))
    assert len(measured["width_samples_px"])<len(plain["width_samples_px"])
    assert "ruler_overlap" in measured["flags"]
    assert measured["width_px"]==pytest.approx(12,abs=.5)


def test_mask_dimensions_and_instances(tmp_path):
    mask=np.zeros((20,30),np.uint16)
    mask[2:6,2:6]=1
    mask[8:12,8:12]=2
    path=tmp_path/"mask.png"
    Image.fromarray(mask).save(path)
    assert set(np.unique(read_labels(path,(20,30,3))))=={0,1,2}
    with pytest.raises(ValueError): read_labels(path,(30,20,3))


def test_exif_orientation(tmp_path):
    image=Image.new("RGB",(60,40),(200,190,180))
    exif=Image.Exif()
    exif[274]=6
    path=tmp_path/"rotated.jpg"
    image.save(path,exif=exif)
    assert read_image(path).shape==(60,40,3)


def test_cli_exports_without_overwriting(tmp_path):
    source=tmp_path/"input"
    source.mkdir()
    masks=tmp_path/"masks"
    masks.mkdir()
    mask,rgb=rod()
    Image.fromarray(rgb).save(source/"rod.png")
    Image.fromarray(mask.astype(np.uint8)*255).save(masks/"rod.png")
    output=tmp_path/"output"
    argv=[str(source),"-o",str(output),"--scale-mode","manual","--um-per-pixel","0.25","--mask-dir",str(masks)]
    assert main(argv)==0
    with (output/"summary.csv").open(encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
    assert len(rows)==1
    assert float(rows[0]["Ширина (мкм)"])==pytest.approx(3,abs=.125)
    assert float(rows[0]["Длина цепи (мкм)"])==pytest.approx(45,abs=.5)
    wb=load_workbook(output/"measurements.xlsx")
    assert wb["Сводная"].max_row==2
    details=list((output/"images").glob("*/details.json"))
    row=json.loads(details[0].read_text())["objects"][0]
    for field in ("crop_original","crop_annotated","object_mask"):
        assert (output/row[field]).is_file()
    # Export the original RGB pixels, not the object labels or a masked image.
    original=np.asarray(Image.open(output/row["crop_original"]))
    x0,y0,x1,y1=row["crop_bbox_xyxy"]
    np.testing.assert_array_equal(original,rgb[y0:y1,x0:x1])
    assert original.shape[2]==3
    index=(output/"index.html").read_text()
    assert row["crop_original"] in index
    assert "Исходный фрагмент фото" in index
    assert list((output/"images").glob("*/edit_mask.html"))
    snapshot=(output/"manifest.json").read_bytes()
    assert main(argv)==2
    assert (output/"manifest.json").read_bytes()==snapshot


def test_unscaled_and_corrupt_images_still_have_report(tmp_path):
    source=tmp_path/"input"
    source.mkdir()
    mask,rgb=rod()
    Image.fromarray(rgb).save(source/"rod.png")
    (source/"bad.jpg").write_bytes(b"invalid image")
    output=tmp_path/"output"
    assert main([str(source),"-o",str(output)])==2
    manifest=json.loads((output/"manifest.json").read_text())
    assert {i["status"] for i in manifest["images"]}=={"error","uncalibrated"}
    with (output/"summary.csv").open(encoding="utf-8-sig",newline="") as f:
        for row in csv.DictReader(f): assert row["Длина цепи (мкм)"]==""


def test_reference_size_mismatch_is_explicit(tmp_path):
    source=tmp_path/"input"
    source.mkdir()
    Image.new("RGB",(100,100),"white").save(source/"photo.png")
    reference=tmp_path/"ref.png"
    Image.new("RGB",(200,200),"white").save(reference)
    output=tmp_path/"output"
    assert main([str(source),"-o",str(output),"--scale-mode","reference","--reference",str(reference),
                 "--reference-points","0","0","100","0","--reference-distance-um","30"])==2
    calibration=json.loads((output/"manifest.json").read_text())["images"][0]["calibration"]
    assert calibration["um_per_px"] is None
    assert "reference_image_dimensions_mismatch" in calibration["reasons"]
