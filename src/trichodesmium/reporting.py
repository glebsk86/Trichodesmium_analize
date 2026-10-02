"""CSV/XLSX exports, overlays and auditable object crops."""
import csv
import json
from pathlib import Path

import cv2
import numpy as np
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SUMMARY_FIELDS = ["Фото", "ID объекта", "Вид", "Ширина (мкм)", "Длина цепи (мкм)",
                  "Длина клетки (мкм)", "Количество клеток", "Статус", "Способ подсчёта"]
DETAIL_FIELDS = ["photo", "object_id", "source_label", "species", "status", "flags",
                 "scale_um_per_px", "calibration_method", "calibration_quality",
                 "calibration_reasons", "measurement_quality", "cell_quality", "species_quality",
                 "length_px", "width_px", "width_std_px", "length_um", "width_um", "width_std_um",
                 "cell_length_px", "cell_length_um", "cell_count", "cell_count_method",
                 "width_sample_count", "cell_interval_count", "bbox_xyxy", "segmentation_method",
                 "ruler_overlap_fraction", "crop_original", "crop_annotated", "object_mask"]


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+"\n",encoding="utf-8")


def flatten(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value,ensure_ascii=False)
    # Keep user-controlled filenames as literal spreadsheet text.
    if isinstance(value,str) and value.startswith(("=","+","-","@")):
        return "'"+value
    return value


def write_csv(path, rows, fields):
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f,fieldnames=fields,extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key:flatten(row.get(key)) for key in fields})


def summary(row):
    return dict(zip(SUMMARY_FIELDS,[row["photo"],row["object_id"],row["species"],row["width_um"],
                row["length_um"],row["cell_length_um"],row["cell_count"],row["status"],row["cell_count_method"]]))


def workbook(path, rows, photos):
    wb = Workbook()
    wb.remove(wb.active)
    for name,data,fields in [("Сводная",[summary(r) for r in rows],SUMMARY_FIELDS),
                             ("Подробная",rows,DETAIL_FIELDS),
                             ("Фотографии",photos,["photo","status","candidate_count","error","calibration"])]:
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in data:
            ws.append([flatten(row.get(k)) for k in fields])
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="275D68")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width=min(48,max(16,len(str(column[0].value))+2))
    note=wb.create_sheet("Прочти сначала")
    for text in ["Экспериментальный прототип. Все объекты требуют проверки.",
                 "Вид не определяется автоматически в версии 0.1.",
                 "Число клеток — оценка по непроверенным кандидатным перегородкам.",
                 "Пустая ячейка означает отсутствие измерения, а не ноль.",
                 "Для обрезанных объектов измеряется только видимый фрагмент.",
                 "CSV с отдельными измерениями ширины/клеток находятся рядом с таблицей фото."]:
        note.append([text])
    note.column_dimensions["A"].width=110
    wb.save(path)


def draw_object(canvas, mask, row):
    contours,_=cv2.findContours(mask.astype(np.uint8),cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
    cv2.drawContours(canvas,contours,-1,(245,100,60),1)
    path=np.round(row["path_xy"]).astype(np.int32)
    if len(path)>1:
        cv2.polylines(canvas,[path],False,(0,220,240),1)
    for point in row["septum_points_xy"]:
        cv2.circle(canvas,tuple(np.round(point).astype(int)),2,(255,20,220),-1)
    for line in row["width_lines_xy"][::max(1,len(row["width_lines_xy"])//20)]:
        a,b=np.round(line).astype(int)
        cv2.line(canvas,tuple(a),tuple(b),(70,220,70),1)
    x,y,_,_=row["bbox_xyxy"]
    cv2.putText(canvas,row["object_id"].split("/")[-1],(x,max(14,y-5)),
                cv2.FONT_HERSHEY_SIMPLEX,.5,(230,50,50),1,cv2.LINE_AA)


def save_visuals(directory,rgb,labels,rows,obstruction,crops=True):
    directory.mkdir(parents=True,exist_ok=True)
    Image.fromarray(labels.astype(np.int32)).save(directory/"labels.tif")
    Image.fromarray(obstruction.astype(np.uint8)*255).save(directory/"ruler_mask.png")
    overlay=rgb.copy()
    if obstruction.any():
        overlay[obstruction]=(.6*overlay[obstruction]+.4*np.array([100,110,255])).astype(np.uint8)
    for row in rows:
        mask=labels==row["source_label"]
        draw_object(overlay,mask,row)
        if not crops: continue
        x0,y0,x1,y1=row["bbox_xyxy"]
        x0,y0=max(0,x0-15),max(0,y0-15)
        x1,y1=min(rgb.shape[1],x1+15),min(rgb.shape[0],y1+15)
        objdir=directory/row["object_id"].split("/")[-1]
        objdir.mkdir(exist_ok=True)
        # Local annotation includes only this object's geometry.
        annotated=rgb.copy()
        draw_object(annotated,mask,row)
        Image.fromarray(rgb[y0:y1,x0:x1]).save(objdir/"original.png")
        Image.fromarray(annotated[y0:y1,x0:x1]).save(objdir/"annotated.png")
        Image.fromarray(mask[y0:y1,x0:x1].astype(np.uint8)*255).save(objdir/"mask.png")
        row["crop_bbox_xyxy"]=[x0,y0,x1,y1]
        for key,file in [("crop_original","original.png"),("crop_annotated","annotated.png"),("object_mask","mask.png")]:
            row[key]=(objdir/file).relative_to(directory.parent.parent).as_posix()
    Image.fromarray(overlay).save(directory/"overview.png")


def per_photo(directory,rows,photo_info):
    write_csv(directory/"details.csv",rows,DETAIL_FIELDS)
    write_json(directory/"details.json",{"image":photo_info,"objects":rows})
    widths=[]
    cells=[]
    for row in rows:
        scale=row["scale_um_per_px"]
        for index,value in enumerate(row["width_samples_px"],1):
            widths.append({"object_id":row["object_id"],"index":index,"width_px":value,
                           "width_um":value*scale if scale else None})
        for index,value in enumerate(row["cell_intervals_px"],1):
            cells.append({"object_id":row["object_id"],"index":index,"length_px":value,
                          "length_um":value*scale if scale else None,"status":"unverified_septum_interval"})
    write_csv(directory/"widths.csv",widths,["object_id","index","width_px","width_um"])
    write_csv(directory/"cell_intervals.csv",cells,["object_id","index","length_px","length_um","status"])
