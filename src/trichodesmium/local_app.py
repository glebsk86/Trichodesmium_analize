"""One local command: photo directory -> proposals -> primary B report."""
import argparse
import csv
from datetime import datetime
import json
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from . import __version__
from . import cli, tube_compare
from .pipeline import run_batch
from .calibration import point_scale,auto_scale
from .imaging import read_image
from .reporting import write_json

COLUMNS = ['Фото','ID кандидата','Статус проверки','Длина поддержанных участков, мкм',
           'Средняя ширина маски, мкм','Длина поддержанных участков, рабочие px',
           'Средняя ширина маски, рабочие px','Сечений ширины','Масштаб исходника, мкм/px',
           'Масштаб рабочего растра X, мкм/px','Масштаб рабочего растра Y, мкм/px',
           'Вид','Полная длина цепи, мкм','Число клеток',
           'Длина прослеженной оси, мкм','Длина прослеженной оси, рабочие px',
           'Невидимая часть оси, мкм','Невидимая часть оси, рабочие px','Тип длины']


def tables(output, manifest):
    rows = []
    for photo in manifest['images']:
        scale = photo['calibration'];factors = scale['working_um_per_px_xy'] or [None,None]
        for proposal in photo['proposals']:
            method = proposal['methods']['ensemble']
            for i,fit in enumerate(method['fits'],1):
                if fit['reasons']:continue
                g = fit['measurement_guides']
                rows.append([photo['photo'],f"{proposal['source_label']}.{i}",
                             tube_compare.STATUS.get(method['status'],method['status']),
                             g['length_um'],g['width_um'],g['visible_axis_length_px'],g['mean_fitted_width_px'],
                             len(g['width_samples_px']),scale['source_calibration']['um_per_px'],
                             *factors,'не определён',None,None,g['traced_length_um'],g['traced_axis_length_px'],
                             g['unobserved_length_um'],g['unobserved_axis_length_px'],g['length_kind']])
    with (output/'B-measurements.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer = csv.writer(stream,delimiter=';');writer.writerow(COLUMNS);writer.writerows(rows)
    wb = Workbook();ws = wb.active;ws.title='Кандидаты B';ws.append(COLUMNS)
    for row in rows:ws.append(row)
    for cell in ws[1]:cell.font = Font(bold=True)
    ws.freeze_panes = 'A2';ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width = min(48,max(16,len(column[0].value or '')+2))
    notes = wb.create_sheet('Определения')
    for line in ["Объекты — непроверенные кандидаты; ID оси не равен доказанному трихому.",
                 "L оси — вся прослеженная линия до концов; разрывы включены как оценка и показаны штрихом.",
                 "L видимого — участки с поддержкой маски; шкала и общие узлы исключены.",
                 "W — средняя ширина подобранной маски по нарисованным сечениям.",
                 "Вид, полная скрытая длина и число клеток не определяются.",
                 "Строки общих маршрутов в узле нельзя суммировать для биомассы."]:
        notes.append([line])
    notes.column_dimensions['A'].width = 105;wb.save(output/'B-measurements.xlsx')
    return len(rows)


def parser():
    p = argparse.ArgumentParser(description='Фото → отчёт B, PNG и таблица измерений (CPU).')
    p.add_argument('input',type=Path,nargs='?',help='Папка JPG/PNG/HEIC/HEIF/TIFF/BMP')
    p.add_argument('-o','--output',type=Path,help='Новая папка результата; по умолчанию results/run-дата-время')
    p.add_argument('--recursive',action='store_true',help='Включать подпапки')
    p.add_argument('--scale-mode',choices=['auto','manual','reference'],default='auto')
    p.add_argument('--um-per-pixel',type=float,help='Ручные мкм на исходный пиксель')
    p.add_argument('--reference',type=Path)
    p.add_argument('--reference-points',type=float,nargs=4)
    p.add_argument('--reference-distance-um',type=float)
    p.add_argument('--mask-dir',type=Path,help='Папка проверенных PNG масок вместо детектора')
    p.add_argument('--mask-format',choices=['auto','binary','instances'],default='auto')
    p.add_argument('--working-width',type=int,default=960,help='Ширина рабочего B-растра; по умолчанию 960 px')
    p.add_argument('--workers',type=int,default=0,help='Linux: число параллельных фото; 0 — авто, 1 — последовательно')
    p.add_argument('--limited-output',action='store_true',help='Только сводный HTML со встроенными картинками и CSV/XLSX')
    p.add_argument('--version',action='version',version=__version__)
    return p


def run(args):
    source = args.input.expanduser().resolve()
    output = (args.output or Path('results')/datetime.now().strftime('run-%Y%m%d-%H%M%S-%f')).expanduser().resolve()
    if output.exists():raise ValueError('Папка результата уже существует; выберите новое имя.')
    if args.working_width<64:raise ValueError('Рабочая ширина должна быть не меньше 64 px.')
    if args.workers<0:raise ValueError('--workers должен быть 0 или положительным числом.')
    seed_args = [str(source),'-o',str(output/'seed-report'),'--scale-mode',args.scale_mode]
    for flag,value in [('--um-per-pixel',args.um_per_pixel),('--reference',args.reference),
                       ('--reference-distance-um',args.reference_distance_um),('--mask-dir',args.mask_dir)]:
        if value is not None:seed_args += [flag,str(value)]
    if args.reference_points:seed_args += ['--reference-points',*[str(v) for v in args.reference_points]]
    if args.recursive:seed_args += ['--recursive']
    seed_args += ['--mask-format',args.mask_format]
    # Validate before creating anything; output inside input must never recurse.
    seed_options = cli.parser().parse_args(seed_args);cli.validate(seed_options)
    args.input=source;args.output=output
    reference_calibration=None;reference_size=None
    if args.scale_mode=='reference':
        ref=read_image(args.reference);reference_size=[ref.shape[1],ref.shape[0]]
        if args.reference_points:
            import numpy as np
            points=np.asarray(args.reference_points).reshape(2,2)
            if np.any(points<0) or np.any(points[:,0]>=ref.shape[1]) or np.any(points[:,1]>=ref.shape[0]):
                raise ValueError('Точки референса вне изображения.')
            reference_calibration=point_scale(args.reference_points,args.reference_distance_um).to_dict()
        else:reference_calibration=auto_scale(ref).to_dict()
        if reference_calibration['um_per_px'] is None:raise ValueError('Масштаб референса не определён; задайте точки.')
    print('Алгоритм B · нормализация фото, объединение фрагментов и измерения',flush=True)
    manifest=run_batch(args,reference_calibration,reference_size)
    compare_code=1 if manifest['errors'] else 0
    count = tables(output,manifest)
    limited=getattr(args,'limited_output',False)
    index = output/'index.html' if limited else output/'report/index.html'
    prefix='' if limited else '../'
    text = index.read_text(encoding='utf-8').replace('<h1>Находки B: фото по порядку</h1>',
            f'<h1>Находки B: фото по порядку</h1><p><a href="{prefix}B-measurements.xlsx">Таблица XLSX</a> · '
            f'<a href="{prefix}B-measurements.csv">Таблица CSV</a></p>')
    index.write_text(text,encoding='utf-8')
    if not limited:
        (output/'index.html').write_text('<!doctype html><meta charset="utf-8"><meta http-equiv="refresh" content="0;url=report/index.html">'
                                       '<a href="report/index.html">Открыть отчёт B</a>',encoding='utf-8')
    info = dict(version=__version__,input=str(source),output=str(output),algorithm="B",
                report_exit_code=compare_code,total_seconds=manifest["total_seconds"],rows=count,summary=manifest['summary'],errors=manifest['errors'])
    if not limited:write_json(output/'run-info.json',info)
    png='' if limited else f'PNG: {output/"report/png-report"}\n'
    print(f'Готово: {len(manifest["images"])} уникальных фото, {count} кандидатных осей.\n'
          f'Откройте: {output/"index.html"}\n{png}'
          f'Таблица: {output/"B-measurements.xlsx"}',flush=True)
    if manifest['errors']:print('Есть ошибки отдельных файлов; см. run-info.json.',file=sys.stderr)
    return 1 if compare_code or not manifest['images'] else 0


def main(argv=None):
    p = parser();args = p.parse_args(argv)
    if args.input is None:
        if not sys.stdin.isatty():p.error('Укажите папку фотографий первым аргументом.')
        try:
            value = input('Путь к папке с фотографиями: ').strip()
        except (EOFError,KeyboardInterrupt):return 2
        if not value:return 2
        args.input = Path(value)
    try:return run(args)
    except (OSError,ValueError) as exc:
        print(f'Ошибка: {exc}',file=sys.stderr);return 2


if __name__=='__main__':
    import multiprocessing
    multiprocessing.freeze_support()
    raise SystemExit(main())
